"""Check that an .xlsx dashboard will keep working when new rows are added.

    python check_dashboard.py dashboard.xlsx [--data-sheet Data] [--dashboard-sheet Dashboard]

Catches the traps that make "paste next month under the last row" silently
break, without opening Excel:

  1. Table column names that do not match the header cells (openpyxl writes
     Column1, Column2... unless you name them). Structured references like
     Sales[Revenue] then break and Excel offers to "repair" the file.
  2. Formula columns in the Data table that are not stored as calculated
     columns, so new rows pasted into the table get no formula.
  3. Dashboard formulas that use fixed ranges ending at today's last row
     (Data!G2:G661), which ignore every row added below.
  4. A dashboard sheet with no formulas at all (typed numbers never update).

Exit code 0 = all checks pass, 1 = at least one FAIL. Read only.
Needs: openpyxl (pip install openpyxl).
"""
import re
import sys
import zipfile

from openpyxl import load_workbook

args = sys.argv[1:]
if not args:
    sys.exit(__doc__)
path = args[0]
data_sheet = args[args.index("--data-sheet") + 1] if "--data-sheet" in args else "Data"
dash_sheet = args[args.index("--dashboard-sheet") + 1] if "--dashboard-sheet" in args else "Dashboard"

fails, warns, oks = [], [], []
wb = load_workbook(path)  # formulas, not cached values
if data_sheet not in wb.sheetnames:
    sys.exit(f"FAIL: no sheet called {data_sheet!r}. Sheets: {wb.sheetnames}")
if dash_sheet not in wb.sheetnames:
    sys.exit(f"FAIL: no sheet called {dash_sheet!r}. Sheets: {wb.sheetnames}")
ws = wb[data_sheet]
last_row = ws.max_row

# ---- 1 and 2: the Data table
tables = list(ws.tables.values())
if not tables:
    warns.append(f"{data_sheet} has no Excel Table. Fine if the dashboard uses whole-column "
                 "references (Data!G:G); otherwise new rows will be missed.")
for t in tables:
    first, last = t.ref.split(":")
    hdr_row = int(re.sub(r"[A-Z]+", "", first))
    header = [c.value for c in ws[t.ref][0]]
    names = [c.name for c in t.tableColumns]
    if not names:
        fails.append(f"Table {t.displayName}: no column names stored (openpyxl default). "
                     "Set tableColumns to the header text.")
    elif names != [str(h) for h in header]:
        bad = [(n, h) for n, h in zip(names, header) if n != str(h)][:3]
        fails.append(f"Table {t.displayName}: column names {bad} do not match the header cells.")
    else:
        oks.append(f"Table {t.displayName}: {len(names)} column names match the headers.")
    if t.ref.split(":")[1] and int(re.sub(r"[A-Z]+", "", last)) != last_row:
        warns.append(f"Table {t.displayName} ends at row {re.sub(r'[A-Z]+', '', last)} "
                     f"but the sheet has data to row {last_row}.")
    # formula columns must be calculated columns
    body = ws[t.ref][1:]
    for i, col in enumerate(t.tableColumns):
        vals = [r[i].value for r in body[:50] if r[i].value is not None]
        has_formula = vals and all(isinstance(v, str) and v.startswith("=") for v in vals)
        if has_formula and col.calculatedColumnFormula is None:
            fails.append(f"Table {t.displayName}, column {col.name!r}: rows hold formulas but the "
                         "column is not a calculated column, so pasted rows get no formula.")
        elif has_formula:
            oks.append(f"Column {col.name!r} is a calculated column.")

# ---- 3 and 4: every other sheet
fixed = re.compile(rf"'?{re.escape(data_sheet)}'?!\$?[A-Z]+\$?(\d+):\$?[A-Z]+\$?(\d+)")
for other in wb.worksheets:
    if other.title == data_sheet:
        continue
    n_formulas = 0
    for row in other.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str) and v.startswith("="):
                n_formulas += 1
                for start, end in fixed.findall(v):
                    if int(end) <= last_row + 5:
                        fails.append(f"{other.title}!{c.coordinate}: range ends at row {end} "
                                     f"(data ends at {last_row}); new rows will be ignored. "
                                     "Use the table name or a whole column.")
                        break
    if n_formulas == 0 and other.title == dash_sheet:
        fails.append(f"{other.title}: no formulas at all. Typed numbers will not update.")
    elif n_formulas:
        oks.append(f"{other.title}: {n_formulas} formulas.")

with zipfile.ZipFile(path) as z:
    charts = [n for n in z.namelist() if n.startswith("xl/charts/chart")]
oks.append(f"{len(charts)} chart(s) in the file.")
if charts:
    warns.append("openpyxl drops charts when it re-saves a file it opened. To add rows, rebuild "
                 "the workbook from the clean data, or edit the sheet XML, never load-and-save.")

for label, items in (("OK", oks), ("WARN", warns), ("FAIL", fails)):
    for m in items:
        print(f"{label:4}  {m}")
print("\nRESULT:", "FAIL" if fails else "PASS")
sys.exit(1 if fails else 0)
